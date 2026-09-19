import os
import time
import json
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client=Groq(api_key=my_api_key)
#model = "llama-3.3-70b-versatile"
model = "openai/gpt-oss-120b"


class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]
    
class SkillGapAnalysis(BaseModel):
    matched_skills: list[str]
    missing_required_skills: list[str]
    missing_preferred_skills: list[str]
    partially_matched_skills: list[str]
    skill_gap_summary: str
    recommendations: list[str]
    
def analyze_skill_gap(job, resume):

    prompt = f"""
    You are an expert technical recruiter.

    Analyze the skill gap between the Job Description and the Candidate Resume.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Compare the candidate's skills, experience, and projects against
    the required and preferred skills in the job description.

    Rules:

   1. matched_skills:
   Include skills that the candidate clearly demonstrates at a sufficient level.
   A skill must NOT also appear in partially_matched_skills.

    2. missing_required_skills:
    Include required skills that are not demonstrated by the candidate.
    A skill must NOT also appear in matched_skills or partially_matched_skills.

    3. missing_preferred_skills:
    Include preferred skills that are not demonstrated by the candidate.

    4. partially_matched_skills:
   Include skills where the candidate demonstrates some knowledge or related
   experience, but not enough evidence to consider the skill fully matched.

   If a required skill is partially demonstrated, classify it only as
   partially_matched_skills and do not also place it in
   missing_required_skills.

   A partially matched skill must NOT also appear in matched_skills
   or missing_required_skills.
   
    5. Each skill or requirement should appear in only ONE category:
   - matched_skills
   - partially_matched_skills
   - missing_required_skills
   - missing_preferred_skills

   Never duplicate the same skill or requirement across categories.
    
    6. skill_gap_summary:
       Give a concise summary of the candidate's main skill gaps.

    7. recommendations:
       Provide practical areas the candidate could improve.

    Do not invent skills or experience.

    Return JSON matching this schema:

    {SkillGapAnalysis.model_json_schema()}
    """

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": "You are an expert technical recruiter."
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={"type": "json_object"}
    )

    data = json.loads(response.choices[0].message.content)

    return SkillGapAnalysis(**data)

jobd_schema = JobD.model_json_schema()

def parse_job_description(job_description):
    system_prompt = f"""
    You are an expert HR assistant.

    Your job is to analyze job descriptions and extract
    structured information from them.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}

    IMPORTANT:
    Do NOT return the schema itself.
    Do NOT return fields like "properties", "title" or "type".
    Fill the schema with actual information extracted from the job description.

    If minimum experience is not mentioned, return null.
    If information for a list is missing, return an empty list.
    Do not invent information.
    """

    user_prompt = f"""
    Analyze the following job description:

    {job_description}
    """

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],
        response_format={"type": "json_object"}
    )

    data = json.loads(response.choices[0].message.content)

    return JobD(**data)



#parse real
class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()
def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



# lets do it now
class InterviewQuestions(BaseModel):
    questions: list[str]

interview_schema = InterviewQuestions.model_json_schema()

def generate_interview_questions(job, resume):

    prompt = f"""
    You are a senior technical interviewer.

    Based on the following Job Description and Candidate Resume,
    generate exactly 5 personalized interview questions.

    The questions should:
    - Focus on the candidate's skills and projects.
    - Test missing or weak skills mentioned in the job description.
    - Be suitable for a technical interview.
    - Do not include answers.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Return JSON matching this schema:

    {interview_schema}
    """

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={"type": "json_object"}
    )

    data = json.loads(response.choices[0].message.content)

    return InterviewQuestions(**data)
